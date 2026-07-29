#!/usr/bin/env python3
"""Audit a resume DOCX or text file for common ATS and delivery risks."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


STANDARD_HEADINGS = {
    "教育背景", "教育经历", "工作经历", "实习经历", "项目经历", "校园经历",
    "项目 / 校园实践", "代表项目", "教育与技能", "技能", "技能与证书",
    "荣誉奖项", "个人总结", "职业摘要",
    "education", "experience", "work experience", "professional experience",
    "projects", "leadership", "skills", "certifications", "awards", "summary",
}

EN_STOP = {
    "and", "the", "for", "with", "that", "this", "from", "will", "your", "you",
    "our", "are", "was", "were", "have", "has", "into", "job", "role", "work",
    "team", "skills", "required", "preferred", "responsibilities", "experience",
}

PLACEHOLDER_RE = re.compile(
    r"\[(?:[^\]]{1,80})\]|<[^>]{1,80}>|待填写|待补充|在此填写|xxx+|tbd|todo",
    re.IGNORECASE,
)
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d[\d ()-]{7,}\d)(?!\d)")
CN_ID_RE = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")


def read_docx(path: Path) -> tuple[str, dict]:
    doc = Document(path)
    body = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            table_text.extend(cell.text.strip() for cell in row.cells if cell.text.strip())

    header_footer = []
    for section in doc.sections:
        header_footer.extend(p.text.strip() for p in section.header.paragraphs if p.text.strip())
        header_footer.extend(p.text.strip() for p in section.footer.paragraphs if p.text.strip())

    font_sizes = []
    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.size:
                font_sizes.append(round(run.font.size.pt, 2))

    has_section_columns = False
    for section in doc.sections:
        cols = section._sectPr.find(qn("w:cols"))
        if cols is not None and int(cols.get(qn("w:num"), "1")) > 1:
            has_section_columns = True

    meta = {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "textboxes": len(doc.part.element.xpath(".//w:txbxContent")),
        "sections": len(doc.sections),
        "multi_column_sections": has_section_columns,
        "header_footer_text": header_footer,
        "minimum_explicit_font_pt": min(font_sizes) if font_sizes else None,
    }
    return "\n".join(body + table_text + header_footer), meta


def read_text(path: Path) -> tuple[str, dict]:
    return path.read_text(encoding="utf-8"), {
        "paragraphs": None,
        "tables": None,
        "inline_shapes": None,
        "textboxes": None,
        "sections": None,
        "multi_column_sections": None,
        "header_footer_text": [],
        "minimum_explicit_font_pt": None,
    }


def auto_keywords(text: str, limit: int = 20) -> list[str]:
    english = [
        token.lower()
        for token in re.findall(r"[A-Za-z][A-Za-z0-9+#./-]{2,}", text)
        if token.lower() not in EN_STOP
    ]
    ascii_ranked = [word for word, _ in Counter(english).most_common(limit)]
    chinese_chunks = re.findall(r"[\u4e00-\u9fff]{2,8}", text)
    cn_ranked = [word for word, _ in Counter(chinese_chunks).most_common(limit)]
    return (ascii_ranked + cn_ranked)[:limit]


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).lower()


def audit(text: str, meta: dict, template_mode: bool, keywords: list[str]) -> dict:
    hard_failures: list[str] = []
    warnings: list[str] = []
    passes: list[str] = []
    text_norm = normalized(text)

    if not template_mode and PLACEHOLDER_RE.search(text):
        hard_failures.append("发现模板占位符或 TODO 文本。")
    elif template_mode:
        passes.append("模板模式：允许占位符。")

    if CN_ID_RE.search(text):
        hard_failures.append("发现疑似中国大陆身份证号码；简历通常不应包含。")

    if not template_mode:
        if EMAIL_RE.search(text):
            passes.append("检测到邮箱。")
        else:
            warnings.append("未检测到邮箱。")
        if PHONE_RE.search(text):
            passes.append("检测到电话号码。")
        else:
            warnings.append("未检测到电话号码。")

    if meta["tables"]:
        warnings.append(f"发现 {meta['tables']} 个表格；若用于布局，可能增加 ATS 解析风险。")
    else:
        passes.append("未发现表格布局。")

    if meta["inline_shapes"]:
        warnings.append(
            f"发现 {meta['inline_shapes']} 个内嵌图形/图片；不要让关键信息只存在于图片中。"
        )
    else:
        passes.append("未发现内嵌图片。")

    if meta["textboxes"]:
        hard_failures.append(f"发现 {meta['textboxes']} 个文本框；ATS 版应改为正文段落。")
    elif meta["textboxes"] == 0:
        passes.append("未发现文本框。")

    if meta["multi_column_sections"]:
        warnings.append("发现多栏分节；检查文本提取顺序，ATS 版优先单栏。")
    elif meta["multi_column_sections"] is False:
        passes.append("文档为单栏分节。")

    if meta["header_footer_text"]:
        warnings.append("页眉/页脚含文本；不要把姓名或联系方式只放在页眉/页脚。")

    min_size = meta["minimum_explicit_font_pt"]
    if min_size is not None and min_size < 9:
        hard_failures.append(f"最小显式字号为 {min_size} pt，低于 9 pt。")
    elif min_size is not None and min_size < 9.5:
        warnings.append(f"最小显式字号为 {min_size} pt，建议确认打印可读性。")

    heading_hits = []
    for line in text.splitlines():
        candidate = line.strip().strip(":：").lower()
        if candidate in STANDARD_HEADINGS:
            heading_hits.append(line.strip())
    if len(set(heading_hits)) >= 3:
        passes.append(f"检测到标准模块标题：{', '.join(dict.fromkeys(heading_hits))}")
    else:
        warnings.append("检测到的标准模块标题少于 3 个；确认标题是否清晰、标准。")

    coverage = {}
    if keywords:
        for keyword in keywords:
            coverage[keyword] = keyword.lower() in text_norm
        missing = [keyword for keyword, present in coverage.items() if not present]
        if missing:
            warnings.append("未覆盖关键词：" + "、".join(missing))
        else:
            passes.append("传入的关键词均已覆盖。")

    status = "fail" if hard_failures else ("warn" if warnings else "pass")
    return {
        "status": status,
        "hard_failures": hard_failures,
        "warnings": warnings,
        "passes": passes,
        "keyword_coverage": coverage,
        "structure": meta,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("resume", type=Path)
    parser.add_argument("--jd", type=Path, help="UTF-8 JD text file for keyword clues.")
    parser.add_argument("--keywords", help="Comma-separated, human-confirmed JD keywords.")
    parser.add_argument("--template-mode", action="store_true")
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()

    if not args.resume.exists():
        parser.error(f"Resume not found: {args.resume}")

    if args.resume.suffix.lower() == ".docx":
        text, meta = read_docx(args.resume)
    else:
        text, meta = read_text(args.resume)

    keywords = []
    if args.keywords:
        keywords.extend(keyword.strip() for keyword in args.keywords.split(",") if keyword.strip())
    if args.jd:
        jd_text = args.jd.read_text(encoding="utf-8")
        keywords.extend(keyword for keyword in auto_keywords(jd_text) if keyword not in keywords)

    report = audit(text, meta, args.template_mode, keywords[:30])
    output = json.dumps(report, ensure_ascii=False, indent=2)
    print(output)
    if args.json_out:
        args.json_out.write_text(output + "\n", encoding="utf-8")

    return {"pass": 0, "warn": 1, "fail": 2}[report["status"]]


if __name__ == "__main__":
    sys.exit(main())
